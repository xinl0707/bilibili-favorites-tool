"""Excel导出模块 - 卡片式布局（含封面图片、头像、详细信息），与原版工具保持一致"""
import json
import math
import os
import string

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Font, Border, Side, PatternFill, colors
from openpyxl.utils.units import pixels_to_EMU

# ========== 图片缩放与定位常量（沿用原版经验值） ==========
conversion_length = 1000 / 26.46

TwoCell_h = 1.02 * conversion_length

# 封面图片大小
MaxWidth_cover = 7.62 * conversion_length
MaxHeight_cover = 4 * conversion_length

# 头像图片大小
MaxWidth_face = 1.69 * conversion_length
MaxHeight_face = 1.02 * conversion_length

alignment_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
alignment_left_top = Alignment(horizontal="left", vertical="top", wrap_text=True)
alignment_left_bottom = Alignment(horizontal="left", vertical="bottom", wrap_text=True)


def filling(start_loc, end_loc, ws):
    """绘制卡片粗边框（参数为左上角/右下角坐标，如 'A1'、'H8'）"""
    x_start = start_loc[0]
    y_start = start_loc[1:]
    x_end = end_loc[0]
    y_end = end_loc[1:]
    len_y = int(y_end) - int(y_start) + 1
    alphabet = string.ascii_uppercase
    len_x = alphabet.index(x_end) - alphabet.index(x_start) + 1
    # 四个角
    ws[start_loc].border = Border(left=Side(style='thick'), top=Side(style='thick'))
    ws[end_loc].border = Border(right=Side(style='thick'), bottom=Side(style='thick'))
    ws[x_end + y_start].border = Border(right=Side(style='thick'), top=Side(style='thick'))
    ws[x_start + y_end].border = Border(left=Side(style='thick'), bottom=Side(style='thick'))
    # 上边
    for i in range(0, len_x - 2):
        ws[alphabet[alphabet.index(x_start) + 1 + i] + y_start].border = Border(top=Side(style='thick'))
    # 下边
    for i in range(0, len_x - 2):
        ws[alphabet[alphabet.index(x_start) + 1 + i] + y_end].border = Border(bottom=Side(style='thick'))
    # 左边
    for i in range(0, len_y - 2):
        ws[x_start + str(int(y_start) + 1 + i)].border = Border(left=Side(style='thick'))
    # 右边
    for i in range(0, len_y - 2):
        ws[x_end + str(int(y_start) + 1 + i)].border = Border(right=Side(style='thick'))


def offset_img(img, row, col):
    """精确设置图片在单元格中的偏移位置（col=0 封面，其他为头像）"""
    p2e = pixels_to_EMU
    h, w = img.height, img.width
    size = XDRPositiveSize2D(p2e(w), p2e(h))

    if col == 0:
        COff = (1 - w / MaxWidth_cover) * 1370000
        if h / MaxHeight_cover > 3 / 4:
            race_r = 4 - h / TwoCell_h
        elif h / MaxHeight_cover > 1 / 2:
            race_r = 3 - h / TwoCell_h
            row += 1
        elif h / MaxHeight_cover > 1 / 4:
            race_r = 2 - h / TwoCell_h
            row += 2
        else:
            race_r = 1 - h / TwoCell_h
            row += 3
    else:
        race_r = 1 - h / TwoCell_h
        COff = (1 - w / MaxWidth_face) * 610000 / 2
    ROff = race_r * 183500
    marker = AnchorMarker(col=col, colOff=COff, row=row, rowOff=ROff)
    img.anchor = OneCellAnchor(_from=marker, ext=size)


def _set_title(ws, title, i):
    ws.merge_cells(start_row=i, start_column=3, end_row=i + 1, end_column=8)
    ss = 'C' + str(i)
    font = Font(name="等线", size=16 if len(title) > 30 else 20)
    ws[ss] = title
    ws[ss].alignment = alignment_center
    ws[ss].font = font


def _set_intro(ws, intro, i):
    ws.merge_cells(start_row=i + 4, start_column=3, end_row=i + 6, end_column=8)
    ss = 'C' + str(i + 4)
    font = Font(name="等线", size=12 if len(intro) > 50 else 14)
    ws[ss] = intro
    ws[ss].alignment = alignment_left_top
    ws[ss].font = font


def _set_cover(ws, img_file, i):
    ws.merge_cells(start_row=i, start_column=1, end_row=i + 7, end_column=2)
    if not os.path.exists(img_file):
        return
    img = Image(img_file)
    if img.width / img.height > MaxWidth_cover / MaxHeight_cover:
        img.height = img.height * MaxWidth_cover / img.width
        img.width = MaxWidth_cover
    else:
        img.width = img.width * MaxHeight_cover / img.height
        img.height = MaxHeight_cover
    offset_img(img, i - 1, 0)
    ws.add_image(img)


def _set_face(ws, img_file, i):
    ws.merge_cells(start_row=i, start_column=9, end_row=i + 1, end_column=9)
    if not os.path.exists(img_file):
        return
    img = Image(img_file)
    if img.width / img.height > MaxWidth_face / MaxHeight_face:
        img.height = img.height * MaxWidth_face / img.width
        img.width = MaxWidth_face
    else:
        img.width = img.width * MaxHeight_face / img.height
        img.height = MaxHeight_face
    offset_img(img, i - 1, 8)
    ws.add_image(img)


def _set_number(ws, i):
    ws.merge_cells(start_row=i + 4, start_column=9, end_row=i + 6, end_column=9)
    ws['I' + str(i + 4)] = math.ceil(i / 10)
    ws['I' + str(i + 4)].font = Font(name="华文行楷", size=36)
    ws['I' + str(i + 4)].alignment = alignment_center


def _mark_deleted(ws, i):
    ws['I' + str(i + 7)].font = Font(color=colors.YELLOW)
    ws['I' + str(i + 7)].fill = PatternFill(patternType='solid', bgColor=colors.BLACK)
    ws['I' + str(i + 7)] = '已失效'


def _set_some(ws, value_list, i):
    """写入播放量/收藏量/弹幕/三个时间/时长/BV/UPid/UP昵称"""
    name_list = ['播放量', '收藏量', '弹幕数量', '上传时间', '发布时间', '收藏时间']
    column_list = ['C', 'D', 'E', 'F', 'G', 'H']
    font = Font(name="等线", size=14)
    for num, (name, column) in enumerate(zip(name_list, column_list)):
        ws[column + str(i + 2)] = name
        ws[column + str(i + 2)].font = font
        ss = column + str(i + 3)
        ws[ss] = value_list[num]
        ws[ss].alignment = alignment_left_bottom
        ws[ss].font = font

    ws['C' + str(i + 7)] = '时长'
    ws['D' + str(i + 7)] = value_list[6]
    ws['E' + str(i + 7)] = 'BV'
    ws['F' + str(i + 7)] = value_list[7]
    ws['G' + str(i + 7)] = 'UPid'
    ws['H' + str(i + 7)] = value_list[8]

    ws.merge_cells(start_row=i + 2, start_column=9, end_row=i + 3, end_column=9)
    ws['I' + str(i + 2)] = value_list[9]
    ws['I' + str(i + 2)].alignment = alignment_center


def _set_column_width(ws):
    for j in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H']:
        ws.column_dimensions[j].width = 20


def _write_video_card(ws, v, count, cover_path, avatar_path):
    """把一个视频写成一个 8 行卡片，count 为卡片起始行"""
    filling('A' + str(count), 'H' + str(count + 7), ws)
    some_list = [
        v['观众数据']['播放量'],
        v['观众数据']['收藏量'],
        v['观众数据']['弹幕数量'],
        v['三个时间']['上传时间'],
        v['三个时间']['发布时间'],
        v['三个时间']['收藏时间'],
        v['视频信息']['时长'],
        v['BV'],
        v['up主']['ID'],
        v['up主']['昵称'],
    ]
    _set_some(ws, some_list, count)
    _set_title(ws, v['视频信息']['标题'], count)
    _set_intro(ws, v['视频信息'].get('简介', ''), count)
    _set_cover(ws, cover_path, count)
    _set_face(ws, avatar_path, count)
    _set_number(ws, count)
    if v.get('是否失效'):
        _mark_deleted(ws, count)


def _sanitize_sheet_name(name):
    """清理 Excel sheet 名的非法字符并截断到 31 字符"""
    for ch in ['[', ']', ':', '*', '?', '/', '\\']:
        name = name.replace(ch, '_')
    return name[:31]


def export_collection_excel(data_dir, output_path, progress_callback=None):
    """导出收藏夹信息 Excel（每个收藏夹一个 sheet，卡片式布局含封面/头像）"""
    info_dir = os.path.join(data_dir, '收藏夹信息')
    files = [f for f in os.listdir(info_dir) if f.endswith('.json')]

    wb = Workbook()
    wb.remove(wb.active)

    cover_root = os.path.join(data_dir, '视频封面')
    avatar_root = os.path.join(data_dir, 'up头像')

    for idx, filename in enumerate(files, 1):
        folder_name = filename.replace('.json', '')
        if progress_callback:
            progress_callback(f'[{idx}/{len(files)}] {folder_name}')

        ws = wb.create_sheet(title=_sanitize_sheet_name(folder_name))
        _set_column_width(ws)

        with open(os.path.join(info_dir, filename), 'r', encoding='utf-8') as f:
            data = json.load(f)

        count = 1
        for v in data.values():
            cover_path = os.path.join(cover_root, folder_name, f'{v["BV"]}.jpg')
            avatar_path = os.path.join(avatar_root, f'{v["up主"]["ID"]}.jpg')
            _write_video_card(ws, v, count, cover_path, avatar_path)
            count += 10

    wb.save(output_path)
    if progress_callback:
        progress_callback(f'导出完成: {output_path}')
    return output_path


def export_study_excel(data_dir, output_path, progress_callback=None):
    """导出学习教程精选 Excel（单 sheet，卡片式布局）"""
    from video_filter import filter_study_videos
    videos = filter_study_videos(data_dir)

    wb = Workbook()
    ws = wb.active
    ws.title = '学习教程精选'
    _set_column_width(ws)

    cover_root = os.path.join(data_dir, '视频封面')
    avatar_root = os.path.join(data_dir, 'up头像')

    count = 1
    for v in videos:
        cover_path = os.path.join(cover_root, v['收藏夹'], f'{v["BV"]}.jpg')
        avatar_path = os.path.join(avatar_root, f'{v["up主"]["ID"]}.jpg')
        _write_video_card(ws, v, count, cover_path, avatar_path)
        count += 10

    wb.save(output_path)
    if progress_callback:
        progress_callback(f'导出完成: {output_path}')
    return output_path


def export_value_excel(data_dir, output_path, progress_callback=None):
    """导出有价值视频推荐 Excel（单 sheet，卡片式布局，Top100）"""
    from video_filter import score_all_videos
    videos = score_all_videos(data_dir)[:100]

    wb = Workbook()
    ws = wb.active
    ws.title = '总榜 Top100'
    _set_column_width(ws)

    cover_root = os.path.join(data_dir, '视频封面')
    avatar_root = os.path.join(data_dir, 'up头像')

    count = 1
    for v in videos:
        cover_path = os.path.join(cover_root, v['收藏夹'], f'{v["BV"]}.jpg')
        avatar_path = os.path.join(avatar_root, f'{v["up主"]["ID"]}.jpg')
        _write_video_card(ws, v, count, cover_path, avatar_path)
        count += 10

    wb.save(output_path)
    if progress_callback:
        progress_callback(f'导出完成: {output_path}')
    return output_path
